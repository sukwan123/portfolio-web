# -*- coding: utf-8 -*-
"""자기소개서를 첨삭용 워드 파일로 굽는다. 내용은 웹 포트폴리오 첫 화면과 같다.

    python3 src/letter_docx.py [출력 경로]

자기소개서 PDF(src/paper.py 의 letter())와 같은 데이터 — 한 줄 소개, 요약,
강점 카드 세 장 — 를 쓴다. 사이트 문안이 바뀌면 다시 돌리기만 하면 된다.
서식은 이력서 워드(src/paper_docx.py)와 같은 규칙을 따른다.
"""
import os
import sys

from docx import Document
from docx.shared import Mm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C                                   # noqa: E402
import paper as P                                     # noqa: E402
from build import months_since                        # noqa: E402
from paper_docx import FONT, GREY, font, h2, hyperlink, para, rich  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(P.OUT, "손석완_자기소개서.docx")
BODY = 10.5


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(21)
    sec.top_margin, sec.bottom_margin = Mm(22), Mm(19)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT)
    st.paragraph_format.line_spacing = 1.45

    # 머리글 — 이름 · 문서명, 연차
    years = months_since(C.SITE["career_start"]) // 12 + 1
    p = para(doc, after=2)
    font(p.add_run(C.SITE["name"]), 20, True)
    font(p.add_run("  자기소개서"), 10.5, False, GREY)
    para(doc, "레벨 디자이너 · %d년 차 (Since %s)" % (years, C.SITE["career_start"][:4]),
         10.5, False, RGBColor(0x3C, 0x43, 0x4B), after=16)

    # 한 줄 소개와 요약
    para(doc, C.HOME["lede"], 11.5, True, after=8)
    rich(para(doc, after=6), C.HOME["sub"].replace("{{career}}", "<b>%s</b>" % P.CAREER), BODY)

    # 강점 세 장
    for _, title, text in C.HOME["feats"]:
        h2(doc, title)
        rich(para(doc, after=4), text, BODY)

    # 맨 아래 — 포트폴리오 링크 하나와 연락처
    p = para(doc, before=18, after=2)
    font(p.add_run("만든 레벨과 플레이 영상은 웹 포트폴리오에 있습니다.  "), 9.5, False, GREY)
    hyperlink(p, P.SITE, P.SITE.replace("https://", ""))
    p = para(doc, after=0)
    font(p.add_run(C.SITE["name"]), 9.5, True)
    font(p.add_run(" · %s · %s" % (C.SITE["email"], C.SITE["tel"])), 9.5, False, GREY)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("자기소개서(워드)", OUT)


if __name__ == "__main__":
    main()
