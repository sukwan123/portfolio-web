# -*- coding: utf-8 -*-
"""이력서를 첨삭용 워드 파일로 굽는다. 내용은 src/paper.py 와 같은 데이터에서 나온다.

    python3 src/paper_docx.py [출력 경로]

서식은 일부러 단순하게 둔다 — 스타일을 걸지 않고 굵은 글씨와 표만 쓴다. 고치다가
서식이 깨질 일이 없어야 한다. 글꼴은 맑은 고딕(윈도우 워드 기본).
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C                                   # noqa: E402
import paper as P                                     # noqa: E402

FONT = "맑은 고딕"
GREY = RGBColor(0x4A, 0x51, 0x59)
LINK = RGBColor(0x6E, 0x4F, 0x27)
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(P.OUT, "손석완_이력서.docx")


def font(run, size=10.5, bold=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def rich(par, html, size=10.5, color=None):
    """<b> 는 굵게, <br> 은 줄바꿈, 나머지 태그는 버린다."""
    html = re.sub(r"<(?!/?b>|br>)[^>]*>", "", str(html)).replace("&amp;", "&")
    for chunk in re.split(r"(<b>.*?</b>|<br>)", html):
        if not chunk:
            continue
        if chunk == "<br>":
            par.add_run().add_break(WD_BREAK.LINE)
        elif chunk.startswith("<b>"):
            font(par.add_run(chunk[3:-4]), size, True, color)
        else:
            font(par.add_run(chunk), size, False, color)


def hyperlink(par, url, text, size=9.5):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attrs in (("w:rFonts", {"w:ascii": FONT, "w:hAnsi": FONT, "w:eastAsia": FONT}),
                       ("w:color", {"w:val": "6E4F27"}), ("w:sz", {"w:val": str(int(size * 2))}),
                       ("w:u", {"w:val": "single"})):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    par._p.append(h)


def para(doc, text="", size=10.5, bold=False, color=None, before=0, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if text:
        font(p.add_run(text), size, bold, color)
    return p


def h2(doc, text):
    p = para(doc, text, 11.5, True, before=14, after=6)
    # 밑줄선 — 문단 아래 테두리
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "2"), ("w:color", "C7CCD3")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    after = {qn(t) for t in ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind",
                            "w:jc", "w:rPr")}
    nxt = next((c for c in ppr if c.tag in after), None)
    if nxt is not None:
        nxt.addprevious(bdr)
    else:
        ppr.append(bdr)
    return p


def kv_table(doc, rows, k_width=32, size=9.5):
    """(항목, 내용) 표. 내용은 <b>·<br> 을 살린다. 내용 자리에 함수를 주면 그걸로 채운다."""
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Mm(k_width), Mm(170 - k_width)
        font(cells[0].paragraphs[0].add_run(k), size, True)
        if callable(v):
            v(cells[1].paragraphs[0])
        else:
            rich(cells[1].paragraphs[0], v, size)
    for row in t.rows:                       # 표 안 문단은 간격을 죄어 둔다
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
    return t


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # 머리글
    p = para(doc, after=2)
    font(p.add_run(C.SITE["name"]), 20, True)
    font(p.add_run("  이력서"), 10.5, False, GREY)
    para(doc, "레벨 디자이너 · %s" % P.CAREER, 10.5, True, after=3)
    p = para(doc, after=2)
    font(p.add_run("%s · %s" % (C.SITE["email"], C.SITE["tel"])), 9.5)
    p = para(doc, after=10)
    font(p.add_run("포트폴리오  "), 9.5, True)
    hyperlink(p, P.SITE + "/index.html", P.SITE.replace("https://", ""))

    # 요약
    h2(doc, "요약")
    para(doc, C.HOME["lede"], 11, True, after=4)
    rich(para(doc, after=6), C.HOME["sub"].replace("{{career}}", "<b>%s</b>" % P.CAREER))
    t = doc.add_table(rows=2, cols=4)
    t.style = "Table Grid"
    for i, (n, lab) in enumerate(C.HOME["stats"]):
        font(t.cell(0, i).paragraphs[0].add_run(P.CAREER if "{{" in n else n), 14, True)
        font(t.cell(1, i).paragraphs[0].add_run(lab), 8.5, False, GREY)

    # 경력 한눈에
    h2(doc, "경력")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, head in enumerate(("기간 · 회사", "프로젝트", "소속 · 직책 / 장르")):
        font(t.rows[0].cells[i].paragraphs[0].add_run(head), 8.5, True, GREY)
    for when, co, title, pos in C.HOME["career"]:
        cells = t.add_row().cells
        font(cells[0].paragraphs[0].add_run(co), 9.5, True)
        cells[0].paragraphs[0].add_run().add_break(WD_BREAK.LINE)
        font(cells[0].paragraphs[0].add_run(when), 8.5, False, GREY)
        font(cells[1].paragraphs[0].add_run(title), 10, True)
        rich(cells[2].paragraphs[0], pos, 9)
    for row in t.rows:
        row.cells[0].width, row.cells[1].width, row.cells[2].width = Mm(38), Mm(52), Mm(80)

    # 경력 상세
    h2(doc, "경력 상세")
    for co, when, title, genre, role, items, ln in P.DETAIL:
        p = para(doc, before=8, after=0)
        font(p.add_run(title), 11.5, True)
        font(p.add_run("   %s · %s" % (co, when)), 8.5, False, GREY)
        para(doc, genre, 8.5, False, GREY, after=0)
        p = para(doc, after=4)
        rich(p, role, 9.5)
        font(p.add_run("  ·  "), 9.5, False, GREY)
        hyperlink(p, "%s/%s" % (P.SITE, ln[0]), ln[1])
        kv_table(doc, items, k_width=26, size=9.5)

    # 보유 기술
    h2(doc, "보유 기술")
    kv_table(doc, C.HOME["skills"])

    # 개인 작업
    h2(doc, "개인 작업")
    rows = []
    for w in C.PERSONAL["own"]:
        if w.get("draft"):
            continue
        def fill(par, w=w):
            rich(par, "<b>%s</b> — %s  " % (w["title"], P.strip(w["meta"])), 9.5)
            hyperlink(par, "%s/%s" % (P.SITE, w["page"]), "웹에서 보기")
        rows.append((w["year"], fill))
    kv_table(doc, rows, k_width=18)

    # 학력 · 기타
    h2(doc, "학력 · 기타")
    kv_table(doc, list(C.HOME["edu"]) + [("업계 외", P.OFFINDUSTRY)])

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("이력서(워드)", OUT)


if __name__ == "__main__":
    main()
